"""
File Parser module for the AI Resume Analyzer.

Handles extraction of raw text from uploaded resume files.
Supports PDF (via PyMuPDF/fitz) and DOCX (via python-docx).
"""

import io
import logging
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

logger = logging.getLogger(__name__)


def extract_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file using PyMuPDF."""
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text_parts = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            page_text = page.get_text("text")
            if page_text.strip():
                text_parts.append(page_text)
        doc.close()
        full_text = "\n".join(text_parts)
        if not full_text.strip():
            raise ValueError("PDF contains no extractable text.")
        return full_text
    except fitz.FitzError as e:
        logger.error(f"PyMuPDF error: {e}")
        raise ValueError(f"Could not read PDF file: {e}")


def extract_from_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file including table content."""
    try:
        doc = Document(io.BytesIO(file_bytes))
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    text_parts.append(row_text)
        full_text = "\n".join(text_parts)
        if not full_text.strip():
            raise ValueError("DOCX contains no extractable text.")
        return full_text
    except Exception as e:
        logger.error(f"Error reading DOCX: {e}")
        raise ValueError(f"Could not read DOCX file: {e}")


def parse_resume(file_bytes: bytes, filename: str) -> str:
    """Route file to correct parser based on extension."""
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        return extract_from_pdf(file_bytes)
    elif extension == ".docx":
        return extract_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: '{extension}'.")
